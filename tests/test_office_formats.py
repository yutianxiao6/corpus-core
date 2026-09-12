from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import ModuleType, SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook
from pptx import Presentation

from corpuscore.contracts.documents import ContentBlockType, LoadedContent
from corpuscore.loaders import ExcelLoader, PowerPointLoader
from corpuscore.parsers import PptxParser, XlsxParser
from corpuscore.parsers.ocr import OcrPdfParser
from corpuscore.sources import FileSystemSourceProvider


def source_for(path: Path):  # type: ignore[no-untyped-def]
    return next(iter(FileSystemSourceProvider([path]).discover()))


class OfficeParserTests(unittest.TestCase):
    def test_xlsx_preserves_sheet_and_header_context(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "inventory.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "库存"
            sheet.append(["产品", "数量"])
            sheet.append(["螺丝", 12])
            workbook.save(path)
            parsed = XlsxParser().parse(ExcelLoader().load(source_for(path)))

        self.assertEqual(parsed.metadata["sheet_count"], 1)
        self.assertEqual(len(parsed.blocks), 2)
        self.assertIn("产品: 螺丝", parsed.blocks[1].content)
        self.assertEqual(parsed.blocks[1].metadata["sheet"], "库存")
        self.assertIs(parsed.blocks[0].block_type, ContentBlockType.TABLE)

    def test_pptx_preserves_slide_text_and_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "briefing.pptx"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "季度报告"
            textbox = slide.shapes.add_textbox(0, 0, 200, 100)
            textbox.text = "收入增长"
            table = slide.shapes.add_table(2, 2, 0, 100, 300, 200).table
            table.cell(0, 0).text = "指标"
            table.cell(0, 1).text = "值"
            table.cell(1, 0).text = "收入"
            table.cell(1, 1).text = "100"
            presentation.save(path)
            parsed = PptxParser().parse(PowerPointLoader().load(source_for(path)))

        self.assertEqual(parsed.title, "季度报告")
        self.assertEqual(parsed.metadata["slide_count"], 1)
        self.assertTrue(any(block.block_type is ContentBlockType.TABLE for block in parsed.blocks))
        self.assertTrue(any("收入增长" in block.content for block in parsed.blocks))

    def test_ocr_parser_emits_image_text_with_page_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "ocr-fixture.pdf"
            path.write_bytes(b"pdf")
            source = source_for(path)
            # Build a content contract directly; OCR itself is tested with deterministic module fakes.
            loaded = LoadedContent(source=source, binary=b"pdf")
            pixmap = SimpleNamespace(width=1, height=1, samples=b"\x00\x00\x00")
            page = SimpleNamespace(get_pixmap=lambda **_: pixmap)

            class FakeDocument:
                def __iter__(self):  # type: ignore[no-untyped-def]
                    return iter((page,))

                def close(self) -> None:
                    return None

            document = FakeDocument()
            fitz = ModuleType("fitz")
            fitz.open = lambda **_: document  # type: ignore[attr-defined]
            pytesseract = ModuleType("pytesseract")
            pytesseract.image_to_string = lambda *_args, **_kwargs: "扫描文本"  # type: ignore[attr-defined]
            with patch.dict("sys.modules", {"fitz": fitz, "pytesseract": pytesseract}):
                parsed = OcrPdfParser(languages="eng", dpi=100).parse(loaded)

        self.assertEqual(parsed.parser_name, "pdf_ocr")
        self.assertEqual(parsed.blocks[0].block_type, ContentBlockType.IMAGE_TEXT)
        self.assertEqual(parsed.blocks[0].page, 1)


if __name__ == "__main__":
    unittest.main()
