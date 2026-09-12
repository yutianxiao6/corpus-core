from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from corpuscore.contracts.documents import ContentBlockType
from corpuscore.exceptions import DocumentLoadError, DocumentParseError
from corpuscore.loaders import BinaryFileLoader, DocxLoader, PdfLoader, TextLoader
from corpuscore.parsers import DocxParser, HtmlParser, PdfParser, StructuredTextParser
from corpuscore.sources import FileSystemSourceProvider


def source_for(path: Path):  # type: ignore[no-untyped-def]
    return next(iter(FileSystemSourceProvider([path]).discover()))


class BinaryLoaderTests(unittest.TestCase):
    def test_binary_loader_verifies_content_hash_and_size(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.pdf"
            path.write_bytes(b"%PDF fixture")
            source = source_for(path)
            loaded = PdfLoader(max_file_size_bytes=20).load(source)
            self.assertEqual(loaded.binary, b"%PDF fixture")
            self.assertIsNone(loaded.text)

            path.write_bytes(b"changed")
            with self.assertRaisesRegex(DocumentLoadError, "changed"):
                PdfLoader().load(source)

    def test_binary_loader_requires_declared_supported_type(self) -> None:
        with self.assertRaises(ValueError):
            BinaryFileLoader(extensions=())


class PdfParserTests(unittest.TestCase):
    def test_pdf_page_numbers_and_title_are_preserved(self) -> None:
        class FakePage:
            def __init__(self, text: str) -> None:
                self.text = text

            def extract_text(self) -> str:
                return self.text

        fake_reader = SimpleNamespace(
            is_encrypted=False,
            pages=[FakePage("第一页内容"), FakePage("Second page")],
            metadata=SimpleNamespace(title="企业手册"),
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.pdf"
            path.write_bytes(b"fixture")
            loaded = PdfLoader().load(source_for(path))
            with patch("corpuscore.parsers.pdf.PdfReader", return_value=fake_reader):
                document = PdfParser().parse(loaded)

        content_blocks = [
            block
            for block in document.blocks
            if block.block_type is not ContentBlockType.PAGE_BREAK
        ]
        self.assertEqual([block.page for block in content_blocks], [1, 2])
        self.assertEqual(document.title, "企业手册")
        self.assertEqual(document.metadata["page_count"], 2)

    def test_image_only_pdf_requests_ocr(self) -> None:
        fake_reader = SimpleNamespace(
            is_encrypted=False,
            pages=[SimpleNamespace(extract_text=lambda: "")],
            metadata=None,
        )
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "scan.pdf"
            path.write_bytes(b"fixture")
            loaded = PdfLoader().load(source_for(path))
            with (
                patch("corpuscore.parsers.pdf.PdfReader", return_value=fake_reader),
                self.assertRaisesRegex(DocumentParseError, "OCR"),
            ):
                PdfParser().parse(loaded)


class DocxParserTests(unittest.TestCase):
    def test_docx_body_order_headings_and_tables(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.docx"
            document = Document()
            document.core_properties.title = "产品文档"
            document.add_heading("安装", level=1)
            document.add_paragraph("Install locally.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "OS"
            table.cell(0, 1).text = "Version"
            table.cell(1, 0).text = "Linux"
            table.cell(1, 1).text = "2026"
            document.save(path)

            parsed = DocxParser().parse(DocxLoader().load(source_for(path)))

        self.assertEqual(parsed.title, "产品文档")
        self.assertEqual(
            [block.block_type for block in parsed.blocks],
            [ContentBlockType.HEADING, ContentBlockType.PARAGRAPH, ContentBlockType.TABLE],
        )
        self.assertIn("Linux | 2026", parsed.blocks[-1].content)


class HtmlParserTests(unittest.TestCase):
    def test_html_ignores_scripts_and_preserves_structure(self) -> None:
        html = """
        <html><head><title>安装手册</title><script>secret()</script></head>
        <body><h1>安装</h1><p>Install <b>offline</b>.</p>
        <pre>def run():\n    return 1</pre>
        <table><tr><th>OS</th><th>Version</th></tr><tr><td>Linux</td><td>2026</td></tr></table>
        </body></html>
        """
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "manual.html"
            path.write_text(html, encoding="utf-8")
            loaded = TextLoader(extensions=(".html",), media_types=("text/html",)).load(
                source_for(path)
            )
            parsed = HtmlParser().parse(loaded)

        self.assertEqual(parsed.title, "安装手册")
        self.assertEqual(
            [block.block_type for block in parsed.blocks],
            [
                ContentBlockType.HEADING,
                ContentBlockType.PARAGRAPH,
                ContentBlockType.CODE,
                ContentBlockType.TABLE,
            ],
        )
        self.assertNotIn("secret", " ".join(block.content for block in parsed.blocks))
        self.assertIn("    return 1", parsed.blocks[2].content)


class StructuredParserTests(unittest.TestCase):
    def _loaded(self, directory: str, filename: str, content: str):  # type: ignore[no-untyped-def]
        path = Path(directory) / filename
        path.write_text(content, encoding="utf-8")
        return TextLoader(extensions=(path.suffix,)).load(source_for(path))

    def test_csv_rows_repeat_header_context(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            parsed = StructuredTextParser("csv").parse(
                self._loaded(directory, "systems.csv", "name,version\nLinux,2026\nWindows,11\n")
            )
        self.assertEqual(len(parsed.blocks), 2)
        self.assertEqual(parsed.blocks[0].content, "name: Linux | version: 2026")
        self.assertEqual(parsed.blocks[0].metadata["headers"], ("name", "version"))

    def test_json_and_jsonl_preserve_record_boundaries(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            json_document = StructuredTextParser("json").parse(
                self._loaded(directory, "records.json", '[{"id":1},{"id":2}]')
            )
            jsonl_document = StructuredTextParser("jsonl").parse(
                self._loaded(directory, "records.jsonl", '{"id":1}\n{"id":2}\n')
            )
        self.assertEqual(len(json_document.blocks), 2)
        self.assertEqual(len(jsonl_document.blocks), 2)
        self.assertEqual(json_document.blocks[0].content, '{"id":1}')


if __name__ == "__main__":
    unittest.main()
