"""PowerPoint parser preserving slide, text and table boundaries."""

from __future__ import annotations

import hashlib
from io import BytesIO

from corpuscore.contracts.documents import (
    ContentBlock,
    ContentBlockType,
    LoadedContent,
    ParsedDocument,
)
from corpuscore.exceptions import DocumentParseError
from corpuscore.normalization import normalize_text


class PptxParser:
    name = "pptx"
    version = "1"

    def parse(self, loaded: LoadedContent) -> ParsedDocument:
        if loaded.binary is None:
            raise DocumentParseError("PowerPoint parser requires binary content")
        try:
            from pptx import Presentation

            presentation = Presentation(BytesIO(loaded.binary))
            blocks: list[ContentBlock] = []
            title: str | None = None
            for slide_number, slide in enumerate(presentation.slides, start=1):
                slide_blocks = self._slide_blocks(slide, slide_number)
                if slide_blocks:
                    blocks.extend(slide_blocks)
                    if title is None:
                        title = next(
                            (
                                block.content
                                for block in slide_blocks
                                if block.block_type is ContentBlockType.HEADING
                            ),
                            None,
                        )
                if slide_number < len(presentation.slides) and slide_blocks:
                    blocks.append(ContentBlock(ContentBlockType.PAGE_BREAK, "", page=slide_number))
        except ImportError as exc:
            raise DocumentParseError(
                "PowerPoint support requires the optional 'office' extra",
                details={"source_id": loaded.source.source_id},
            ) from exc
        except Exception as exc:
            raise DocumentParseError(
                "cannot parse PowerPoint source", details={"source_id": loaded.source.source_id}
            ) from exc
        if not any(block.block_type is not ContentBlockType.PAGE_BREAK for block in blocks):
            raise DocumentParseError(
                "PowerPoint presentation has no indexable content",
                details={"source_id": loaded.source.source_id},
            )
        material = (
            f"{loaded.source.source_id}\0{loaded.source.content_hash}\0{self.name}\0{self.version}"
        )
        return ParsedDocument(
            document_id=hashlib.sha256(material.encode()).hexdigest(),
            source=loaded.source,
            title=title,
            blocks=blocks,
            parser_name=self.name,
            parser_version=self.version,
            metadata={"format": "pptx", "slide_count": len(presentation.slides)},
        )

    @staticmethod
    def _slide_blocks(slide: object, slide_number: int) -> list[ContentBlock]:
        blocks: list[ContentBlock] = []
        for shape in slide.shapes:  # type: ignore[attr-defined]
            if getattr(shape, "has_table", False):
                rows = [
                    " | ".join(normalize_text(cell.text) for cell in row.cells)
                    for row in shape.table.rows
                ]
                content = "\n".join(row for row in rows if row)
                if content:
                    blocks.append(
                        ContentBlock(
                            ContentBlockType.TABLE,
                            content,
                            page=slide_number,
                            metadata={"slide": slide_number},
                        )
                    )
            elif getattr(shape, "has_text_frame", False):
                content = normalize_text(getattr(shape, "text", ""))
                if not content:
                    continue
                is_title = "title" in str(getattr(shape, "name", "")).lower()
                blocks.append(
                    ContentBlock(
                        ContentBlockType.HEADING if is_title else ContentBlockType.PARAGRAPH,
                        content,
                        page=slide_number,
                    )
                )
        return blocks
