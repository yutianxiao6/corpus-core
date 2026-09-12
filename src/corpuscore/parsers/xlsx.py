"""Excel parser preserving worksheet and row boundaries."""

from __future__ import annotations

import hashlib
from io import BytesIO

from corpuscore.contracts.documents import (
    ContentBlock,
    ContentBlockType,
    LoadedContent,
    ParsedDocument,
)
from corpuscore.exceptions import DocumentParseError


class XlsxParser:
    name = "xlsx"
    version = "1"

    def parse(self, loaded: LoadedContent) -> ParsedDocument:
        if loaded.binary is None:
            raise DocumentParseError("Excel parser requires binary content")
        try:
            from openpyxl import load_workbook  # type: ignore[import-untyped]

            workbook = load_workbook(BytesIO(loaded.binary), read_only=True, data_only=True)
            blocks: list[ContentBlock] = []
            sheet_count = 0
            try:
                for sheet in workbook.worksheets:
                    sheet_count += 1
                    rows = list(sheet.iter_rows(values_only=True))
                    headers = self._values(rows[0]) if rows else ()
                    for row_number, row in enumerate(rows, start=1):
                        values = self._values(row)
                        if not any(values):
                            continue
                        if headers and row_number > 1:
                            content = " | ".join(
                                f"{headers[index] if index < len(headers) else f'column_{index + 1}'}: "
                                f"{value}"
                                for index, value in enumerate(values)
                            )
                        else:
                            content = " | ".join(values)
                        blocks.append(
                            ContentBlock(
                                ContentBlockType.TABLE,
                                content,
                                metadata={
                                    "sheet": sheet.title,
                                    "row_number": row_number,
                                    "headers": headers,
                                },
                            )
                        )
            finally:
                workbook.close()
        except DocumentParseError:
            raise
        except ImportError as exc:
            raise DocumentParseError(
                "Excel support requires the optional 'tables' extra",
                details={"source_id": loaded.source.source_id},
            ) from exc
        except Exception as exc:
            raise DocumentParseError(
                "cannot parse Excel source", details={"source_id": loaded.source.source_id}
            ) from exc
        if not blocks:
            raise DocumentParseError(
                "Excel workbook has no indexable rows",
                details={"source_id": loaded.source.source_id},
            )
        material = (
            f"{loaded.source.source_id}\0{loaded.source.content_hash}\0{self.name}\0{self.version}"
        )
        return ParsedDocument(
            document_id=hashlib.sha256(material.encode()).hexdigest(),
            source=loaded.source,
            blocks=blocks,
            parser_name=self.name,
            parser_version=self.version,
            metadata={"format": "xlsx", "sheet_count": sheet_count},
        )

    @staticmethod
    def _values(row: tuple[object, ...] | list[object]) -> tuple[str, ...]:
        return tuple("" if value is None else str(value).strip() for value in row)
